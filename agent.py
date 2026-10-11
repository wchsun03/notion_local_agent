import os
import re
import csv
import json
import sqlite3
import requests

from pathlib import Path
from datetime import datetime, timezone

from dotenv import load_dotenv
from ollama import chat
from prompt_toolkit import prompt
from prompt_toolkit.history import InMemoryHistory

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


# ============================================================
# Generic Local Mac + Notion Agent - V5
# ============================================================

load_dotenv()

NOTION_TOKEN = os.getenv("NOTION_TOKEN")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "qwen3:8b")

NOTION_VERSION = "2026-03-11"
STATE_DB = "state.db"

if not NOTION_TOKEN:
    raise ValueError("找不到 NOTION_TOKEN，請檢查 .env")

NOTION_HEADERS = {
    "Authorization": f"Bearer {NOTION_TOKEN}",
    "Content-Type": "application/json",
    "Notion-Version": NOTION_VERSION,
}


# ============================================================
# Local Mac Security
# ============================================================

HOME = Path.home().resolve()

ALLOWED_ROOTS = [
    (HOME / "Desktop").resolve(),
    (HOME / "Documents").resolve(),
    (HOME / "Desktop" / "AgentWorkspace").resolve(),
]

BLOCKED_NAMES = {
    ".env",
    ".ssh",
    ".aws",
    ".config",
    "id_rsa",
    "id_ed25519",
}

BLOCKED_SUFFIXES = {
    ".pem",
    ".key",
    ".p12",
    ".pfx",
}

SUPPORTED_FILE_TYPES = {
    ".txt",
    ".md",
    ".csv",
    ".xlsx",
    ".docx",
    ".pdf",
}


def is_path_allowed(path: Path) -> bool:

    path = path.expanduser().resolve()

    allowed = any(
        path == root or root in path.parents
        for root in ALLOWED_ROOTS
    )

    if not allowed:
        return False

    for part in path.parts:
        if part in BLOCKED_NAMES:
            return False

    if path.name in BLOCKED_NAMES:
        return False

    if path.suffix.lower() in BLOCKED_SUFFIXES:
        return False

    return True


def resolve_local_path(path_string: str) -> Path:

    raw = Path(path_string).expanduser()

    if raw.is_absolute():
        candidate = raw.resolve()

        if not is_path_allowed(candidate):
            raise PermissionError(
                f"不允許存取此路徑：{candidate}"
            )

        return candidate

    # Try relative to current working directory
    candidate = Path.cwd() / raw

    if candidate.exists():

        candidate = candidate.resolve()

        if is_path_allowed(candidate):
            return candidate

    # Try allowed roots
    matches = []

    for root in ALLOWED_ROOTS:

        candidate = root / raw

        if candidate.exists():
            matches.append(
                candidate.resolve()
            )

    unique = list(
        dict.fromkeys(matches)
    )

    if len(unique) == 1:
        return unique[0]

    if len(unique) > 1:
        raise ValueError(
            f"找到多個同名路徑：{path_string}"
        )

    raise FileNotFoundError(
        f"找不到本機路徑：{path_string}"
    )


# ============================================================
# SQLite Persistent State
# ============================================================

def init_state_db():

    conn = sqlite3.connect(STATE_DB)

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS state (
            key TEXT PRIMARY KEY,
            value TEXT
        )
        """
    )

    conn.commit()
    conn.close()


def set_state(key, value):

    conn = sqlite3.connect(STATE_DB)

    conn.execute(
        """
        INSERT INTO state(key, value)
        VALUES (?, ?)
        ON CONFLICT(key)
        DO UPDATE SET value=excluded.value
        """,
        (
            key,
            json.dumps(
                value,
                ensure_ascii=False
            )
        )
    )

    conn.commit()
    conn.close()


def get_state(key, default=None):

    conn = sqlite3.connect(STATE_DB)

    cursor = conn.execute(
        """
        SELECT value
        FROM state
        WHERE key=?
        """,
        (key,)
    )

    row = cursor.fetchone()

    conn.close()

    if not row:
        return default

    try:
        return json.loads(row[0])
    except Exception:
        return default


def delete_state(key):

    conn = sqlite3.connect(STATE_DB)

    conn.execute(
        "DELETE FROM state WHERE key=?",
        (key,)
    )

    conn.commit()
    conn.close()


def clear_state():

    conn = sqlite3.connect(STATE_DB)

    conn.execute(
        "DELETE FROM state"
    )

    conn.commit()
    conn.close()


def current_state_summary():

    return {
        "last_page":
            get_state("last_page"),

        "last_database":
            get_state("last_database"),

        "last_record":
            get_state("last_record"),

        "last_file":
            get_state("last_file"),

        "pending_batch":
            get_state("pending_batch"),

        "pending_import":
            get_state("pending_import"),
    }


init_state_db()


# ============================================================
# General Helpers
# ============================================================

def rich_text(content):

    return [
        {
            "type": "text",
            "text": {
                "content": str(content)
            }
        }
    ]


def extract_plain_text(items):

    return "".join(
        item.get("plain_text", "")
        for item in (items or [])
    )


def extract_page_title(page):

    properties = page.get(
        "properties",
        {}
    )

    for _, prop in properties.items():

        if prop.get("type") == "title":

            return extract_plain_text(
                prop.get(
                    "title",
                    []
                )
            )

    return ""


def extract_data_source_title(data_source):

    title = extract_plain_text(
        data_source.get(
            "title",
            []
        )
    )

    if title:
        return title

    return data_source.get("name") or ""


def normalize_date(value):

    if value is None:
        return None

    value = str(value).strip()

    formats = [
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
    ]

    for fmt in formats:

        try:

            return datetime.strptime(
                value,
                fmt
            ).strftime(
                "%Y-%m-%d"
            )

        except ValueError:
            pass

    if re.match(
        r"^\d{4}-\d{2}-\d{2}T",
        value
    ):
        return value

    raise ValueError(
        f"無法辨識日期格式：{value}"
    )


def notion_error(response):

    try:

        data = response.json()

        return (
            data.get("message")
            or json.dumps(
                data,
                ensure_ascii=False
            )
        )

    except Exception:

        return response.text


def sanitize_output(text):

    if not text:
        return ""

    text = re.sub(
        r"\[([^\]]+)\]\((https?://[^\)]+)\)",
        r"\1: \2",
        text
    )

    text = re.sub(
        r"\*\*(.*?)\*\*",
        r"\1",
        text
    )

    return text.strip()


# ============================================================
# Local File Tools
# ============================================================

def list_directory(path: str) -> str:

    target = resolve_local_path(path)

    if not target.is_dir():

        return json.dumps(
            {
                "success": False,
                "error": "指定路徑不是資料夾。"
            },
            ensure_ascii=False
        )

    items = []

    for item in sorted(
        target.iterdir(),
        key=lambda x: (
            not x.is_dir(),
            x.name.lower()
        )
    ):

        if not is_path_allowed(item):
            continue

        items.append(
            {
                "name":
                    item.name,

                "type":
                    (
                        "directory"
                        if item.is_dir()
                        else "file"
                    ),

                "path":
                    str(item),

                "extension":
                    (
                        item.suffix.lower()
                        if item.is_file()
                        else None
                    )
            }
        )

    set_state(
        "last_file",
        {
            "path": str(target),
            "type": "directory"
        }
    )

    return json.dumps(
        {
            "success": True,
            "path": str(target),
            "count": len(items),
            "items": items
        },
        ensure_ascii=False
    )


def search_local_files(
    query: str,
    base_path: str = "~/Desktop"
) -> str:

    base = resolve_local_path(
        base_path
    )

    if not base.is_dir():

        return json.dumps(
            {
                "success": False,
                "error": "base_path 不是資料夾。"
            },
            ensure_ascii=False
        )

    query_lower = (
        query
        .strip()
        .lower()
    )

    matches = []

    for root, dirs, files in os.walk(base):

        root_path = Path(root)

        dirs[:] = [
            d
            for d in dirs
            if is_path_allowed(
                root_path / d
            )
        ]

        for filename in files:

            path = (
                root_path
                / filename
            )

            if not is_path_allowed(path):
                continue

            if (
                query_lower
                in filename.lower()
            ):

                matches.append(
                    {
                        "name":
                            filename,

                        "path":
                            str(
                                path.resolve()
                            ),

                        "extension":
                            path.suffix.lower()
                    }
                )

        if len(matches) >= 100:
            break

    return json.dumps(
        {
            "success": True,
            "query": query,
            "count": len(matches),
            "matches": matches[:100]
        },
        ensure_ascii=False
    )


def read_text_file(path: Path):

    return path.read_text(
        encoding="utf-8",
        errors="replace"
    )


def read_pdf_file(path: Path):

    reader = PdfReader(
        str(path)
    )

    pages = []

    for index, page in enumerate(
        reader.pages
    ):

        pages.append(
            {
                "page":
                    index + 1,

                "text":
                    (
                        page.extract_text()
                        or ""
                    )
            }
        )

    return {
        "page_count":
            len(pages),

        "pages":
            pages
    }


def read_docx_file(path: Path):

    document = Document(
        str(path)
    )

    paragraphs = [
        paragraph.text
        for paragraph
        in document.paragraphs
        if paragraph.text.strip()
    ]

    tables = []

    for table in document.tables:

        rows = []

        for row in table.rows:

            rows.append(
                [
                    cell.text
                    for cell
                    in row.cells
                ]
            )

        tables.append(rows)

    return {
        "paragraphs":
            paragraphs,

        "tables":
            tables
    }


def read_csv_file(
    path: Path,
    max_rows=200
):

    with path.open(
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline=""
    ) as file:

        reader = csv.reader(file)

        rows = []

        for index, row in enumerate(reader):

            if index >= max_rows:
                break

            rows.append(row)

    return {
        "rows":
            rows,

        "row_count_returned":
            len(rows),

        "truncated":
            (
                len(rows)
                >= max_rows
            )
    }


def read_xlsx_file(
    path: Path,
    sheet_name=None,
    max_rows=200
):

    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True
    )

    if sheet_name:

        if sheet_name not in workbook.sheetnames:

            raise ValueError(
                (
                    f"找不到工作表："
                    f"{sheet_name}"
                )
            )

        sheets = [
            workbook[
                sheet_name
            ]
        ]

    else:

        sheets = [
            workbook[name]
            for name
            in workbook.sheetnames
        ]

    output = {}

    for sheet in sheets:

        rows = []

        for row_index, row in enumerate(
            sheet.iter_rows(
                values_only=True
            )
        ):

            if row_index >= max_rows:
                break

            rows.append(
                [
                    value
                    for value in row
                ]
            )

        output[
            sheet.title
        ] = {
            "rows":
                rows,

            "row_count_returned":
                len(rows),

            "truncated":
                (
                    len(rows)
                    >= max_rows
                )
        }

    workbook.close()

    return {
        "sheet_names":
            list(output.keys()),

        "sheets":
            output
    }


def read_local_file(
    path: str,
    sheet_name: str | None = None,
    max_rows: int = 200
) -> str:

    target = resolve_local_path(
        path
    )

    if not target.is_file():

        return json.dumps(
            {
                "success": False,
                "error": "指定路徑不是檔案。"
            },
            ensure_ascii=False
        )

    extension = (
        target.suffix
        .lower()
    )

    if extension not in SUPPORTED_FILE_TYPES:

        return json.dumps(
            {
                "success": False,

                "error":
                    (
                        "目前不支援此檔案格式："
                        f"{extension}"
                    )
            },
            ensure_ascii=False
        )

    try:

        if extension in {
            ".txt",
            ".md"
        }:

            content = (
                read_text_file(
                    target
                )
            )

            result = {
                "text":
                    content[:50000],

                "truncated":
                    len(content) > 50000
            }

        elif extension == ".pdf":

            result = read_pdf_file(
                target
            )

        elif extension == ".docx":

            result = read_docx_file(
                target
            )

        elif extension == ".csv":

            result = read_csv_file(
                target,
                max_rows=max_rows
            )

        elif extension == ".xlsx":

            result = read_xlsx_file(
                target,
                sheet_name=
                    sheet_name,
                max_rows=
                    max_rows
            )

        else:

            raise ValueError(
                "Unsupported file type."
            )

    except Exception as exc:

        return json.dumps(
            {
                "success": False,
                "error": str(exc)
            },
            ensure_ascii=False
        )

    set_state(
        "last_file",
        {
            "path":
                str(target),

            "type":
                "file",

            "extension":
                extension,
        }
    )

    return json.dumps(
        {
            "success": True,

            "path":
                str(target),

            "extension":
                extension,

            "content":
                result
        },
        ensure_ascii=False,
        default=str
    )


# ============================================================
# Notion Page Search
# ============================================================

def search_pages(page_name):

    response = requests.post(
        "https://api.notion.com/v1/search",

        headers=NOTION_HEADERS,

        json={
            "query":
                page_name,

            "filter": {
                "property":
                    "object",

                "value":
                    "page"
            },

            "page_size":
                100
        },

        timeout=30
    )

    response.raise_for_status()

    results = []

    for page in (
        response
        .json()
        .get(
            "results",
            []
        )
    ):

        results.append(
            {
                "id":
                    page.get("id"),

                "title":
                    extract_page_title(
                        page
                    ),

                "url":
                    page.get("url"),

                "raw":
                    page,
            }
        )

    return results


def resolve_page(page_name):

    results = search_pages(
        page_name
    )

    target = (
        page_name
        .strip()
        .lower()
    )

    exact = [
        item
        for item in results
        if (
            item["title"]
            .strip()
            .lower()
            == target
        )
    ]

    if len(exact) == 1:

        return {
            "success": True,
            "page": exact[0]
        }

    partial = [
        item
        for item in results

        if (
            target
            in item["title"]
            .strip()
            .lower()

            or

            item["title"]
            .strip()
            .lower()
            in target
        )
    ]

    if len(partial) == 1:

        return {
            "success": True,
            "page": partial[0]
        }

    if len(exact) > 1 or len(partial) > 1:

        return {
            "success": False,

            "error":
                "找到多個相似頁面，請提供更完整名稱。"
        }

    return {
        "success": False,
        "error":
            f"找不到頁面：{page_name}"
    }


# ============================================================
# Notion Page Tools
# ============================================================

def get_page(
    page_name: str
) -> str:

    result = resolve_page(
        page_name
    )

    if not result[
        "success"
    ]:

        return json.dumps(
            result,
            ensure_ascii=False
        )

    page = result[
        "page"
    ]

    set_state(
        "last_page",
        {
            "title":
                page["title"],

            "page_id":
                page["id"],

            "url":
                page["url"],
        }
    )

    return json.dumps(
        {
            "success": True,
            "title":
                page["title"],
            "page_id":
                page["id"],
            "url":
                page["url"],
        },
        ensure_ascii=False
    )


def create_page_under_parent(
    parent_page_id,
    title
):

    response = requests.post(
        "https://api.notion.com/v1/pages",

        headers=NOTION_HEADERS,

        json={
            "parent": {
                "type":
                    "page_id",

                "page_id":
                    parent_page_id
            },

            "properties": {
                "title": {
                    "type":
                        "title",

                    "title":
                        rich_text(
                            title
                        )
                }
            }
        },

        timeout=30
    )

    response.raise_for_status()

    page = response.json()

    return {
        "title":
            (
                extract_page_title(
                    page
                )
                or title
            ),

        "page_id":
            page.get("id"),

        "url":
            page.get("url"),
    }


def create_child_pages(
    parent_page_name: str,
    child_page_titles: list[str]
) -> str:

    resolved = resolve_page(
        parent_page_name
    )

    if not resolved[
        "success"
    ]:

        return json.dumps(
            resolved,
            ensure_ascii=False
        )

    parent = resolved[
        "page"
    ]

    created = []
    errors = []

    for title in child_page_titles:

        try:

            page = (
                create_page_under_parent(
                    parent["id"],
                    title
                )
            )

            created.append(
                page
            )

        except Exception as exc:

            errors.append(
                {
                    "title":
                        title,

                    "error":
                        str(exc)
                }
            )

    return json.dumps(
        {
            "success":
                len(errors) == 0,

            "parent":
                parent["title"],

            "created":
                created,

            "errors":
                errors,
        },

        ensure_ascii=False
    )


# ============================================================
# Notion Blocks
# ============================================================

def append_blocks_to_page(
    page_name: str,
    blocks: list[dict]
) -> str:

    resolved = resolve_page(
        page_name
    )

    if not resolved[
        "success"
    ]:

        return json.dumps(
            resolved,
            ensure_ascii=False
        )

    page = resolved[
        "page"
    ]

    children = []

    for block in blocks:

        block_type = (
            block.get("type")
        )

        text = (
            block.get("text")
        )

        if (
            not block_type
            or text is None
            or not str(text).strip()
        ):

            return json.dumps(
                {
                    "success": False,
                    "error":
                        "Block 缺少 type 或 text。"
                },
                ensure_ascii=False
            )

        rt = rich_text(
            str(text).strip()
        )

        if block_type in {
            "heading_1",
            "heading_2",
            "heading_3",
            "paragraph",
            "bulleted_list_item",
            "numbered_list_item",
        }:

            children.append(
                {
                    "object":
                        "block",

                    "type":
                        block_type,

                    block_type: {
                        "rich_text":
                            rt
                    }
                }
            )

        elif block_type == "to_do":

            children.append(
                {
                    "object":
                        "block",

                    "type":
                        "to_do",

                    "to_do": {
                        "rich_text":
                            rt,

                        "checked":
                            bool(
                                block.get(
                                    "checked",
                                    False
                                )
                            )
                    }
                }
            )

    response = requests.patch(
        (
            "https://api.notion.com/v1/"
            f"blocks/{page['id']}/children"
        ),

        headers=NOTION_HEADERS,

        json={
            "children":
                children
        },

        timeout=30
    )

    response.raise_for_status()

    return json.dumps(
        {
            "success": True,
            "page":
                page["title"],
            "block_count":
                len(children)
        },
        ensure_ascii=False
    )


# ============================================================
# Data Source / Database
# ============================================================

def search_data_sources(
    database_name
):

    response = requests.post(
        "https://api.notion.com/v1/search",

        headers=NOTION_HEADERS,

        json={
            "query":
                database_name,

            "filter": {
                "property":
                    "object",

                "value":
                    "data_source"
            },

            "page_size":
                100
        },

        timeout=30
    )

    response.raise_for_status()

    results = []

    for item in (
        response.json()
        .get(
            "results",
            []
        )
    ):

        results.append(
            {
                "id":
                    item.get("id"),

                "title":
                    extract_data_source_title(
                        item
                    ),

                "url":
                    item.get("url"),

                "raw":
                    item,
            }
        )

    return results


def resolve_data_source(
    database_name
):

    results = search_data_sources(
        database_name
    )

    target = (
        database_name
        .strip()
        .lower()
    )

    exact = [
        item
        for item in results
        if (
            item["title"]
            .strip()
            .lower()
            == target
        )
    ]

    if len(exact) == 1:

        return {
            "success": True,
            "data_source":
                exact[0]
        }

    partial = [
        item
        for item in results

        if (
            target
            in item["title"]
            .strip()
            .lower()

            or

            item["title"]
            .strip()
            .lower()
            in target
        )
    ]

    if len(partial) == 1:

        return {
            "success": True,
            "data_source":
                partial[0]
        }

    return {
        "success": False,

        "error":
            (
                "找不到或無法唯一辨識 Database："
                f"{database_name}"
            )
    }


def retrieve_data_source(
    data_source_id
):

    response = requests.get(
        (
            "https://api.notion.com/v1/"
            f"data_sources/{data_source_id}"
        ),

        headers=NOTION_HEADERS,

        timeout=30
    )

    response.raise_for_status()

    return response.json()


def schema_summary(properties):

    output = {}

    for name, prop in properties.items():

        prop_type = (
            prop.get("type")
        )

        output[name] = {
            "type":
                prop_type
        }

        config = prop.get(
            prop_type,
            {}
        )

        if (
            prop_type
            in {
                "select",
                "multi_select",
                "status"
            }

            and isinstance(
                config,
                dict
            )
        ):

            output[name][
                "options"
            ] = [
                option.get("name")
                for option
                in config.get(
                    "options",
                    []
                )
            ]

    return output


def get_database_schema(
    database_name: str
) -> str:

    resolved = resolve_data_source(
        database_name
    )

    if not resolved[
        "success"
    ]:

        return json.dumps(
            resolved,
            ensure_ascii=False
        )

    ds = resolved[
        "data_source"
    ]

    full = retrieve_data_source(
        ds["id"]
    )

    set_state(
        "last_database",
        {
            "title":
                ds["title"],

            "data_source_id":
                ds["id"],
        }
    )

    return json.dumps(
        {
            "success":
                True,

            "database":
                ds["title"],

            "data_source_id":
                ds["id"],

            "properties":
                schema_summary(
                    full.get(
                        "properties",
                        {}
                    )
                )
        },

        ensure_ascii=False
    )


# ============================================================
# Property Encoding
# ============================================================

def resolve_property_name(
    schema,
    requested_name
):

    requested = (
        requested_name
        .strip()
        .lower()
    )

    exact = [
        name
        for name
        in schema.keys()

        if (
            name.strip().lower()
            == requested
        )
    ]

    if len(exact) == 1:
        return exact[0]

    raise ValueError(
        f"找不到欄位：{requested_name}"
    )


def encode_property_value(
    prop_schema,
    value
):

    prop_type = (
        prop_schema.get(
            "type"
        )
    )

    if prop_type == "title":

        return {
            "title":
                rich_text(value)
        }

    if prop_type == "rich_text":

        return {
            "rich_text":
                rich_text(value)
        }

    if prop_type == "number":

        if isinstance(
            value,
            str
        ):

            value = float(
                value
                .replace(",", "")
            )

        return {
            "number":
                value
        }

    if prop_type == "select":

        return {
            "select": {
                "name":
                    str(value)
            }
        }

    if prop_type == "status":

        return {
            "status": {
                "name":
                    str(value)
            }
        }

    if prop_type == "multi_select":

        if isinstance(
            value,
            str
        ):

            value = [
                item.strip()
                for item
                in value.split(",")
                if item.strip()
            ]

        return {
            "multi_select": [
                {
                    "name":
                        str(item)
                }
                for item
                in value
            ]
        }

    if prop_type == "date":

        return {
            "date": {
                "start":
                    normalize_date(
                        value
                    )
            }
        }

    if prop_type == "checkbox":

        return {
            "checkbox":
                bool(value)
        }

    if prop_type == "url":

        return {
            "url":
                str(value)
        }

    if prop_type == "email":

        return {
            "email":
                str(value)
        }

    if prop_type == "phone_number":

        return {
            "phone_number":
                str(value)
        }

    raise ValueError(
        (
            "目前無法寫入欄位型別："
            f"{prop_type}"
        )
    )


def build_page_properties(
    schema,
    values
):

    output = {}

    for (
        requested_name,
        value
    ) in values.items():

        if value is None:
            continue

        actual_name = (
            resolve_property_name(
                schema,
                requested_name
            )
        )

        output[
            actual_name
        ] = (
            encode_property_value(
                schema[actual_name],
                value
            )
        )

    return output


# ============================================================
# Record Parsing
# ============================================================

def extract_property_value(prop):

    prop_type = (
        prop.get("type")
    )

    if prop_type == "title":

        return extract_plain_text(
            prop.get(
                "title",
                []
            )
        )

    if prop_type == "rich_text":

        return extract_plain_text(
            prop.get(
                "rich_text",
                []
            )
        )

    if prop_type == "number":

        return prop.get(
            "number"
        )

    if prop_type in {
        "select",
        "status"
    }:

        value = prop.get(
            prop_type
        )

        return (
            value.get("name")
            if value
            else None
        )

    if prop_type == "multi_select":

        return [
            item.get("name")
            for item
            in prop.get(
                "multi_select",
                []
            )
        ]

    if prop_type == "date":

        value = prop.get(
            "date"
        )

        return (
            value.get("start")
            if value
            else None
        )

    if prop_type == "checkbox":

        return prop.get(
            "checkbox"
        )

    if prop_type == "url":

        return prop.get(
            "url"
        )

    if prop_type == "email":

        return prop.get(
            "email"
        )

    if prop_type == "phone_number":

        return prop.get(
            "phone_number"
        )

    return None


def normalize_record(page):

    output = {
        "page_id":
            page.get("id"),

        "url":
            page.get("url"),
    }

    for name, prop in (
        page
        .get(
            "properties",
            {}
        )
        .items()
    ):

        output[name] = (
            extract_property_value(
                prop
            )
        )

    return output


def find_title_property(schema):

    for name, prop in schema.items():

        if (
            prop.get("type")
            == "title"
        ):
            return name

    return None


# ============================================================
# Record Create
# ============================================================

def create_database_record(
    database_name: str,
    values: dict
) -> str:

    resolved = resolve_data_source(
        database_name
    )

    if not resolved[
        "success"
    ]:

        return json.dumps(
            resolved,
            ensure_ascii=False
        )

    ds = resolved[
        "data_source"
    ]

    full_ds = retrieve_data_source(
        ds["id"]
    )

    schema = full_ds.get(
        "properties",
        {}
    )

    try:

        properties = (
            build_page_properties(
                schema,
                values
            )
        )

    except Exception as exc:

        return json.dumps(
            {
                "success": False,
                "error": str(exc),
                "available_properties":
                    schema_summary(
                        schema
                    )
            },

            ensure_ascii=False
        )

    response = requests.post(
        "https://api.notion.com/v1/pages",

        headers=NOTION_HEADERS,

        json={
            "parent": {
                "type":
                    "data_source_id",

                "data_source_id":
                    ds["id"]
            },

            "properties":
                properties
        },

        timeout=30
    )

    if response.status_code not in {
        200,
        201
    }:

        return json.dumps(
            {
                "success":
                    False,

                "error":
                    notion_error(
                        response
                    )
            },

            ensure_ascii=False
        )

    created = response.json()

    record = normalize_record(
        created
    )

    title_name = find_title_property(
        schema
    )

    set_state(
        "last_database",
        {
            "title":
                ds["title"],

            "data_source_id":
                ds["id"],
        }
    )

    set_state(
        "last_record",
        {
            "database":
                ds["title"],

            "title":
                (
                    record.get(
                        title_name
                    )
                    if title_name
                    else None
                ),

            "page_id":
                record.get(
                    "page_id"
                ),

            "url":
                record.get(
                    "url"
                ),
        }
    )

    return json.dumps(
        {
            "success": True,
            "database":
                ds["title"],
            "created_record":
                record
        },
        ensure_ascii=False
    )


# ============================================================
# Record Query
# ============================================================

def build_query_filter(
    schema,
    condition
):

    requested_name = (
        condition.get(
            "property"
        )
    )

    operator = (
        condition.get(
            "operator",
            "equals"
        )
    )

    value = condition.get(
        "value"
    )

    property_name = (
        resolve_property_name(
            schema,
            requested_name
        )
    )

    prop_type = (
        schema[
            property_name
        ]
        .get("type")
    )

    if prop_type == "date":

        if operator not in {
            "is_empty",
            "is_not_empty"
        }:

            value = normalize_date(
                value
            )

    return {
        "property":
            property_name,

        prop_type: {
            operator:
                (
                    True
                    if operator
                    in {
                        "is_empty",
                        "is_not_empty"
                    }
                    else value
                )
        }
    }


def build_filter_group(
    schema,
    filters,
    logic="and"
):

    items = [
        build_query_filter(
            schema,
            condition
        )
        for condition
        in filters
    ]

    if len(items) == 1:
        return items[0]

    return {
        logic:
            items
    }


def query_database_records(
    database_name: str,
    filters: list[dict],
    logic: str = "and"
) -> str:

    resolved = resolve_data_source(
        database_name
    )

    if not resolved[
        "success"
    ]:

        return json.dumps(
            resolved,
            ensure_ascii=False
        )

    ds = resolved[
        "data_source"
    ]

    full_ds = retrieve_data_source(
        ds["id"]
    )

    schema = full_ds.get(
        "properties",
        {}
    )

    try:

        notion_filter = (
            build_filter_group(
                schema,
                filters,
                logic
            )
        )

    except Exception as exc:

        return json.dumps(
            {
                "success":
                    False,

                "error":
                    str(exc)
            },

            ensure_ascii=False
        )

    response = requests.post(
        (
            "https://api.notion.com/v1/"
            f"data_sources/{ds['id']}/query"
        ),

        headers=NOTION_HEADERS,

        json={
            "filter":
                notion_filter,

            "page_size":
                100
        },

        timeout=30
    )

    response.raise_for_status()

    records = [
        normalize_record(
            page
        )
        for page
        in response.json().get(
            "results",
            []
        )
    ]

    return json.dumps(
        {
            "success": True,
            "database":
                ds["title"],
            "count":
                len(records),
            "records":
                records
        },
        ensure_ascii=False
    )


# ============================================================
# Batch Update Preview
# ============================================================

def preview_batch_update(
    database_name: str,
    filters: list[dict],
    values: dict,
    logic: str = "and"
) -> str:

    query_result = json.loads(
        query_database_records(
            database_name,
            filters,
            logic
        )
    )

    if not query_result.get(
        "success"
    ):

        return json.dumps(
            query_result,
            ensure_ascii=False
        )

    records = query_result[
        "records"
    ]

    if not records:

        return json.dumps(
            {
                "success": False,
                "error":
                    "沒有符合條件的資料。"
            },
            ensure_ascii=False
        )

    pending = {
        "database":
            database_name,

        "filters":
            filters,

        "logic":
            logic,

        "values":
            values,

        "records":
            records,

        "created_at":
            datetime.now(
                timezone.utc
            ).isoformat(),
    }

    set_state(
        "pending_batch",
        pending
    )

    return json.dumps(
        {
            "success": True,
            "executed": False,
            "confirmation_required":
                True,

            "count":
                len(records),

            "records":
                records,

            "planned_changes":
                values
        },
        ensure_ascii=False
    )


def update_record_by_page_id(
    page_id,
    schema,
    values
):

    properties = (
        build_page_properties(
            schema,
            values
        )
    )

    response = requests.patch(
        (
            "https://api.notion.com/v1/"
            f"pages/{page_id}"
        ),

        headers=NOTION_HEADERS,

        json={
            "properties":
                properties
        },

        timeout=30
    )

    response.raise_for_status()

    return normalize_record(
        response.json()
    )


def confirm_pending_batch() -> str:

    pending = get_state(
        "pending_batch"
    )

    if not pending:

        return json.dumps(
            {
                "success": False,
                "error":
                    "沒有待確認批次操作。"
            },
            ensure_ascii=False
        )

    resolved = resolve_data_source(
        pending[
            "database"
        ]
    )

    if not resolved[
        "success"
    ]:

        return json.dumps(
            resolved,
            ensure_ascii=False
        )

    ds = resolved[
        "data_source"
    ]

    schema = retrieve_data_source(
        ds["id"]
    ).get(
        "properties",
        {}
    )

    updated = []
    errors = []

    for record in pending[
        "records"
    ]:

        try:

            updated.append(
                update_record_by_page_id(
                    record[
                        "page_id"
                    ],
                    schema,
                    pending[
                        "values"
                    ]
                )
            )

        except Exception as exc:

            errors.append(
                str(exc)
            )

    delete_state(
        "pending_batch"
    )

    return json.dumps(
        {
            "success":
                len(errors) == 0,

            "updated_count":
                len(updated),

            "failed_count":
                len(errors),

            "updated":
                updated,

            "errors":
                errors
        },
        ensure_ascii=False
    )


def cancel_pending_batch():

    delete_state(
        "pending_batch"
    )

    return json.dumps(
        {
            "success": True,
            "cancelled": True
        },
        ensure_ascii=False
    )


# ============================================================
# Spreadsheet Import Preview
# ============================================================

def normalize_header(
    value
):

    if value is None:
        return ""

    return (
        str(value)
        .strip()
        .lower()
        .replace(" ", "")
        .replace("_", "")
        .replace("-", "")
    )


def prepare_spreadsheet_import(
    file_path: str,
    database_name: str,
    sheet_name: str | None = None
) -> str:

    target = resolve_local_path(
        file_path
    )

    extension = (
        target.suffix.lower()
    )

    if extension not in {
        ".xlsx",
        ".csv"
    }:

        return json.dumps(
            {
                "success": False,

                "error":
                    (
                        "匯入目前只支援 "
                        "CSV 或 XLSX。"
                    )
            },
            ensure_ascii=False
        )

    resolved = resolve_data_source(
        database_name
    )

    if not resolved[
        "success"
    ]:

        return json.dumps(
            resolved,
            ensure_ascii=False
        )

    ds = resolved[
        "data_source"
    ]

    schema = retrieve_data_source(
        ds["id"]
    ).get(
        "properties",
        {}
    )

    if extension == ".xlsx":

        workbook = load_workbook(
            str(target),
            read_only=True,
            data_only=True
        )

        if sheet_name:

            if sheet_name not in workbook.sheetnames:

                workbook.close()

                return json.dumps(
                    {
                        "success": False,

                        "error":
                            (
                                "找不到工作表："
                                f"{sheet_name}"
                            )
                    },
                    ensure_ascii=False
                )

            sheet = workbook[
                sheet_name
            ]

        else:

            sheet = workbook[
                workbook.sheetnames[0]
            ]

        rows = list(
            sheet.iter_rows(
                values_only=True
            )
        )

        actual_sheet_name = (
            sheet.title
        )

        workbook.close()

    else:

        with target.open(
            "r",
            encoding="utf-8-sig",
            errors="replace",
            newline=""
        ) as file:

            rows = list(
                csv.reader(file)
            )

        actual_sheet_name = None

    if len(rows) < 2:

        return json.dumps(
            {
                "success": False,
                "error":
                    "檔案沒有可匯入資料。"
            },
            ensure_ascii=False
        )

    headers = [
        str(value).strip()
        if value is not None
        else ""
        for value
        in rows[0]
    ]

    schema_map = {
        normalize_header(name):
            name
        for name
        in schema.keys()
    }

    column_mapping = {}

    unmapped_headers = []

    for index, header in enumerate(
        headers
    ):

        normalized = (
            normalize_header(
                header
            )
        )

        if (
            normalized
            in schema_map
        ):

            column_mapping[
                index
            ] = (
                schema_map[
                    normalized
                ]
            )

        elif header:

            unmapped_headers.append(
                header
            )

    prepared_rows = []

    for raw_row in rows[1:]:

        values = {}

        for index, notion_name in (
            column_mapping.items()
        ):

            if index >= len(raw_row):
                continue

            value = raw_row[
                index
            ]

            if (
                value is None
                or str(value).strip()
                == ""
            ):
                continue

            values[
                notion_name
            ] = value

        if values:
            prepared_rows.append(
                values
            )

    pending = {
        "file_path":
            str(target),

        "sheet_name":
            actual_sheet_name,

        "database":
            ds["title"],

        "data_source_id":
            ds["id"],

        "mapping":
            {
                headers[index]:
                    notion_name

                for index, notion_name
                in column_mapping.items()
            },

        "unmapped_headers":
            unmapped_headers,

        "rows":
            prepared_rows,

        "count":
            len(prepared_rows),

        "created_at":
            datetime.now(
                timezone.utc
            ).isoformat(),
    }

    set_state(
        "pending_import",
        pending
    )

    return json.dumps(
        {
            "success": True,

            "executed": False,

            "confirmation_required":
                True,

            "file":
                str(target),

            "sheet":
                actual_sheet_name,

            "database":
                ds["title"],

            "row_count":
                len(
                    prepared_rows
                ),

            "mapping":
                pending[
                    "mapping"
                ],

            "unmapped_headers":
                unmapped_headers,

            "preview_rows":
                prepared_rows[:5]
        },

        ensure_ascii=False,
        default=str
    )


def confirm_pending_import() -> str:

    pending = get_state(
        "pending_import"
    )

    if not pending:

        return json.dumps(
            {
                "success": False,
                "error":
                    "目前沒有待確認匯入。"
            },
            ensure_ascii=False
        )

    ds_id = pending[
        "data_source_id"
    ]

    schema = retrieve_data_source(
        ds_id
    ).get(
        "properties",
        {}
    )

    created = []
    errors = []

    for row_number, values in enumerate(
        pending["rows"],
        start=2
    ):

        try:

            properties = (
                build_page_properties(
                    schema,
                    values
                )
            )

            response = requests.post(
                "https://api.notion.com/v1/pages",

                headers=NOTION_HEADERS,

                json={
                    "parent": {
                        "type":
                            "data_source_id",

                        "data_source_id":
                            ds_id
                    },

                    "properties":
                        properties
                },

                timeout=30
            )

            if response.status_code not in {
                200,
                201
            }:

                errors.append(
                    {
                        "row":
                            row_number,

                        "error":
                            notion_error(
                                response
                            )
                    }
                )

                continue

            created.append(
                normalize_record(
                    response.json()
                )
            )

        except Exception as exc:

            errors.append(
                {
                    "row":
                        row_number,

                    "error":
                        str(exc)
                }
            )

    delete_state(
        "pending_import"
    )

    return json.dumps(
        {
            "success":
                len(errors) == 0,

            "created_count":
                len(created),

            "failed_count":
                len(errors),

            "created":
                created,

            "errors":
                errors
        },
        ensure_ascii=False
    )


def cancel_pending_import():

    delete_state(
        "pending_import"
    )

    return json.dumps(
        {
            "success": True,
            "cancelled": True
        },
        ensure_ascii=False
    )


# ============================================================
# Tools
# ============================================================

TOOLS = [

    # Local Mac
    list_directory,
    search_local_files,

    {
        "type": "function",
        "function": {
            "name":
                "read_local_file",

            "description":
                (
                    "Read a supported local Mac file. "
                    "Supports txt, md, csv, xlsx, docx and pdf."
                ),

            "parameters": {
                "type":
                    "object",

                "properties": {
                    "path": {
                        "type":
                            "string"
                    },

                    "sheet_name": {
                        "type":
                            "string"
                    },

                    "max_rows": {
                        "type":
                            "integer"
                    }
                },

                "required": [
                    "path"
                ]
            }
        }
    },

    # Notion
    get_page,
    create_child_pages,

    {
        "type": "function",
        "function": {
            "name":
                "append_blocks_to_page",

            "parameters": {
                "type":
                    "object",

                "properties": {
                    "page_name": {
                        "type":
                            "string"
                    },

                    "blocks": {
                        "type":
                            "array",

                        "items": {
                            "type":
                                "object",

                            "properties": {
                                "type": {
                                    "type":
                                        "string",

                                    "enum": [
                                        "heading_1",
                                        "heading_2",
                                        "heading_3",
                                        "paragraph",
                                        "bulleted_list_item",
                                        "numbered_list_item",
                                        "to_do"
                                    ]
                                },

                                "text": {
                                    "type":
                                        "string"
                                },

                                "checked": {
                                    "type":
                                        "boolean"
                                }
                            },

                            "required": [
                                "type",
                                "text"
                            ]
                        }
                    }
                },

                "required": [
                    "page_name",
                    "blocks"
                ]
            }
        }
    },

    get_database_schema,

    {
        "type": "function",
        "function": {
            "name":
                "create_database_record",

            "parameters": {
                "type":
                    "object",

                "properties": {
                    "database_name": {
                        "type":
                            "string"
                    },

                    "values": {
                        "type":
                            "object",

                        "additionalProperties":
                            True
                    }
                },

                "required": [
                    "database_name",
                    "values"
                ]
            }
        }
    },

    {
        "type": "function",
        "function": {
            "name":
                "query_database_records",

            "parameters": {
                "type":
                    "object",

                "properties": {
                    "database_name": {
                        "type":
                            "string"
                    },

                    "filters": {
                        "type":
                            "array",

                        "items": {
                            "type":
                                "object",

                            "properties": {
                                "property": {
                                    "type":
                                        "string"
                                },

                                "operator": {
                                    "type":
                                        "string"
                                },

                                "value": {}
                            },

                            "required": [
                                "property"
                            ]
                        }
                    },

                    "logic": {
                        "type":
                            "string",

                        "enum": [
                            "and",
                            "or"
                        ]
                    }
                },

                "required": [
                    "database_name",
                    "filters"
                ]
            }
        }
    },

    {
        "type": "function",
        "function": {
            "name":
                "preview_batch_update",

            "parameters": {
                "type":
                    "object",

                "properties": {
                    "database_name": {
                        "type":
                            "string"
                    },

                    "filters": {
                        "type":
                            "array",

                        "items": {
                            "type":
                                "object"
                        }
                    },

                    "values": {
                        "type":
                            "object",

                        "additionalProperties":
                            True
                    },

                    "logic": {
                        "type":
                            "string",

                        "enum": [
                            "and",
                            "or"
                        ]
                    }
                },

                "required": [
                    "database_name",
                    "filters",
                    "values"
                ]
            }
        }
    },

    {
        "type": "function",
        "function": {
            "name":
                "confirm_pending_batch",

            "parameters": {
                "type":
                    "object",
                "properties": {}
            }
        }
    },

    {
        "type": "function",
        "function": {
            "name":
                "cancel_pending_batch",

            "parameters": {
                "type":
                    "object",
                "properties": {}
            }
        }
    },

    {
        "type": "function",
        "function": {
            "name":
                "prepare_spreadsheet_import",

            "description":
                (
                    "Prepare an XLSX or CSV import into "
                    "a Notion database. "
                    "This only previews and requires confirmation."
                ),

            "parameters": {
                "type":
                    "object",

                "properties": {
                    "file_path": {
                        "type":
                            "string"
                    },

                    "database_name": {
                        "type":
                            "string"
                    },

                    "sheet_name": {
                        "type":
                            "string"
                    }
                },

                "required": [
                    "file_path",
                    "database_name"
                ]
            }
        }
    },

    {
        "type": "function",
        "function": {
            "name":
                "confirm_pending_import",

            "parameters": {
                "type":
                    "object",
                "properties": {}
            }
        }
    },

    {
        "type": "function",
        "function": {
            "name":
                "cancel_pending_import",

            "parameters": {
                "type":
                    "object",
                "properties": {}
            }
        }
    },
]


FUNCTION_MAP = {

    "list_directory":
        list_directory,

    "search_local_files":
        search_local_files,

    "read_local_file":
        read_local_file,

    "get_page":
        get_page,

    "create_child_pages":
        create_child_pages,

    "append_blocks_to_page":
        append_blocks_to_page,

    "get_database_schema":
        get_database_schema,

    "create_database_record":
        create_database_record,

    "query_database_records":
        query_database_records,

    "preview_batch_update":
        preview_batch_update,

    "confirm_pending_batch":
        confirm_pending_batch,

    "cancel_pending_batch":
        cancel_pending_batch,

    "prepare_spreadsheet_import":
        prepare_spreadsheet_import,

    "confirm_pending_import":
        confirm_pending_import,

    "cancel_pending_import":
        cancel_pending_import,
}


# ============================================================
# System Prompt
# ============================================================

SYSTEM_PROMPT = """
你是一個部署在 Mac 上的本機 AI Agent。

你可以：

1. 讀取 Mac 本機允許範圍內的檔案。
2. 搜尋與列出 Desktop / Documents 的檔案。
3. 讀取 TXT / MD / CSV / XLSX / DOCX / PDF。
4. 操作 Notion Page。
5. 操作 Notion Database。
6. 查詢與新增 Notion records。
7. 執行批次更新 preview。
8. 執行 Spreadsheet -> Notion import preview。

本機檔案安全規則：

1.
本機檔案預設只能 READ。

不得：
- 刪除
- 移動
- 改名
- 覆寫
- 修改

除非未來新增專門的 write tool。

2.
不得嘗試繞過 Python path allowlist。

3.
如果使用者只說「桌面」，
優先理解為 ~/Desktop。

4.
如果使用者說「剛才那個檔案」，
可以根據 conversation history 或 persistent last_file 判斷。

Spreadsheet Import 規則：

1.
CSV/XLSX 匯入 Notion 前，
必須先使用 prepare_spreadsheet_import。

2.
Preview 後不得自行確認。

3.
只有使用者明確說：
「確認」
「確認匯入」
「執行匯入」

才使用 confirm_pending_import。

4.
若有 unmapped headers，
必須向使用者說明哪些欄位不會匯入。

Batch Update 規則：

批次修改 Notion 前必須 preview。

只有使用者確認後才真正執行。

Notion 規則：

1.
需要操作 Notion 時必須使用工具。

2.
只有 Tool Result success=true
才能宣稱成功。

3.
不得捏造資料。

4.
多筆候選時不得自行猜測。

回覆風格：

- 使用繁體中文
- 不使用 emoji
- 不使用 Markdown 粗體
- 不使用 Markdown hyperlink
- 簡潔、專業、直接
- 不描述內部思考過程
"""


# ============================================================
# Conversation
# ============================================================

def new_conversation():

    return [
        {
            "role":
                "system",

            "content":
                (
                    SYSTEM_PROMPT
                    +
                    "\n\nPersistent state:\n"
                    +
                    json.dumps(
                        current_state_summary(),
                        ensure_ascii=False
                    )
                )
        }
    ]


# ============================================================
# Agent Runner
# ============================================================

def run_agent(
    user_input,
    messages
):

    messages.append(
        {
            "role":
                "user",

            "content":
                user_input
        }
    )

    for _ in range(15):

        response = chat(
            model=
                OLLAMA_MODEL,

            messages=
                messages,

            tools=
                TOOLS,

            options={
                "temperature":
                    0.15
            }
        )

        messages.append(
            response.message
        )

        tool_calls = (
            response
            .message
            .tool_calls
            or []
        )

        if not tool_calls:

            return sanitize_output(
                response
                .message
                .content
            )

        for call in tool_calls:

            function_name = (
                call
                .function
                .name
            )

            arguments = (
                call
                .function
                .arguments
            )

            function = (
                FUNCTION_MAP.get(
                    function_name
                )
            )

            if not function:

                result = json.dumps(
                    {
                        "success":
                            False,

                        "error":
                            (
                                "未知工具："
                                f"{function_name}"
                            )
                    },

                    ensure_ascii=False
                )

            else:

                try:

                    result = function(
                        **arguments
                    )

                except Exception as exc:

                    result = json.dumps(
                        {
                            "success":
                                False,

                            "error":
                                str(exc)
                        },

                        ensure_ascii=False
                    )

            messages.append(
                {
                    "role":
                        "tool",

                    "tool_name":
                        function_name,

                    "content":
                        result,
                }
            )

    return (
        "Agent 達到最大工具呼叫次數，"
        "已停止。"
    )


# ============================================================
# Terminal UI
# ============================================================

if __name__ == "__main__":

    print()
    print("=" * 64)
    print(
        "Local Mac + Notion Agent - V5"
    )
    print("=" * 64)

    print(
        "Local Files READ / Notion READ+WRITE / Import Preview"
    )

    print()

    print(
        "/state          顯示 persistent state"
    )

    print(
        "/clear          清除本次對話"
    )

    print(
        "/forget         清除 persistent state"
    )

    print(
        "/cancel         取消 pending batch"
    )

    print(
        "/cancel-import  取消 pending import"
    )

    print(
        "exit            離開"
    )

    print()

    history = (
        InMemoryHistory()
    )

    messages = (
        new_conversation()
    )

    while True:

        user_input = prompt(
            "You > ",
            history=history
        ).strip()

        command = (
            user_input.lower()
        )

        if command in {
            "exit",
            "quit",
            "q"
        }:

            print(
                "Agent 已結束。"
            )

            break

        if command == "/state":

            print()

            print(
                json.dumps(
                    current_state_summary(),
                    ensure_ascii=False,
                    indent=2
                )
            )

            print()

            continue

        if command == "/clear":

            messages = (
                new_conversation()
            )

            print(
                "\n本次對話已清除。\n"
            )

            continue

        if command == "/forget":

            clear_state()

            messages = (
                new_conversation()
            )

            print(
                "\nPersistent state 已清除。\n"
            )

            continue

        if command == "/cancel":

            print()

            print(
                cancel_pending_batch()
            )

            print()

            continue

        if command == "/cancel-import":

            print()

            print(
                cancel_pending_import()
            )

            print()

            continue

        if not user_input:
            continue

        try:

            result = run_agent(
                user_input,
                messages
            )

            print()
            print("Agent >")
            print(result)
            print()

        except KeyboardInterrupt:

            print(
                "\nAgent 已中止。"
            )

            break

        except Exception as exc:

            print()
            print(
                "ERROR:",
                exc
            )
            print()

